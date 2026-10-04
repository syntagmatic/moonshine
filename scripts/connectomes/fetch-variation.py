# /// script
# dependencies = ["openpyxl"]
# ///
"""Development, sex and neuropeptide data for connectomes article 8 (Same Genes, Different Wiring).

Run: uv run scripts/connectomes/fetch-variation.py

Sources, cached under scripts/connectomes/.cache/worm/ (gitignored):

Development: Witvliet, Mulcahy, Mitchell et al. 2021, "Connectomes across development reveal
principles of brain maturation", Nature 596: 257-261. Eight isogenic hermaphrodites, the brain
(nerve ring and ventral ganglion) only. Edge lists from NemaNode's download-connectivity API
(datasets witvliet_2020_1 to _8, the same content as WormWiring's xlsx files); ages from
NemaNode /api/datasets ("time", hours after birth: 0, 5, 8, 16, 23, 27, 50, 50); the class of
every chemical connection from the paper's Supplementary Table 6 (MOESM8). Cell classes and
types from NemaNode /api/cells.

Sexes: Cook, Jarrell, Brittin et al. 2019, "Whole-animal connectomes of both Caenorhabditis
elegans sexes", Nature 571: 63-71. Weights are EM serial sections of connectivity (synapse
number and size), not synapse counts. From the Nature supplement:
- Supplementary Information 8 (MOESM12): hermaphrodite left/right comparisons, the left and
  right member of a homologous pair onto a class that is itself a left/right pair (chemical
  1,285 rows, gap junctions 705), with Cook's sd and Z.
- Supplementary Information 9 (MOESM13): male against hermaphrodite for cell classes (chemical
  1,823 rows of which two are exact duplicates, so 1,821 pairs; gap junctions from both direction
  sheets, since the gap-junction 'all pairs' sheet repeats only the hermaphrodite-stronger one, less
  three pairs listed in both with contradictory weights: 525 pairs), with the number of cells in the presynaptic class, Cook's
  variance estimates and Z; the per-direction sheets mark differences confirmed by fluorescence
  ('this work', and 'Oren-Suisse', the sheet's spelling of Oren-Suissa, Bayer & Hobert 2016) and those involving dimorphic neurons.
And from WormWiring: SI 5 corrected July 2020 (the hermaphrodite chemical matrix, 302 neurons,
used against the neuropeptide network).

Neuropeptides: Ripoll-Sanchez, Watteyne, Sun et al. 2023, "The neuropeptidergic connectome of
C. elegans", Neuron 111: 3570-3589 (CC BY). Aggregate matrices from github.com/
LidiaRipollSanchez/Neuropeptide-Connectome (MIT), pinned to commit 6689619, the 2024-02-02
update (dmsr-5 and npr-34): short-, mid- and long-range models, 302 x 302, rows sending,
weight = number of neuropeptide-receptor couples linking the pair (CeNGEN threshold 4, EC50
500 nM). Neuron groups from the repository's 072022_anatomical_class.csv.

Writes, in docs/connectomes/shared/data/:
worm-witvliet.json
  ages, stages: per dataset
  classes: the five MOESM8 labels; nodes: [{name, cls, role}] (role s/i/m from NemaNode, b muscle,
  o other); chem: [[pre, post, class, [synapses in datasets 1..8]]]
worm-sexes.json
  lr.chem, lr.gap: [[A, B, left, right, cookSd]]  (A, B cell classes)
  sex.chem, sex.gap: [[A, B, herm, male, cells in A, cookZ, flag]]
  flag: 'c' confirmed by Cook et al., 'o' by Oren-Suissa et al. 2016, 'd' involves a dimorphic
  neuron (PHA, PHB, PHC, LUA, AVL, DVB, ...), '' otherwise
worm-peptide.json
  nodes: [{name, cls, role, group}] in the neuropeptide matrices' order (their Fig. 5);
  groups: anatomical class names; syn: [[pre, post, sections]] (Cook hermaphrodite chemical);
  pep: {short, mid, long}: 302 x 302 weights row-major, one base-36 character per pair
"""
import csv, io, json, re, sys, urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / 'scripts/connectomes/.cache/worm'
OUT = ROOT / 'docs/connectomes/shared/data'
NATURE = 'https://static-content.springer.com/esm/art%3A10.1038%2F{}/MediaObjects/{}'
NP = 'https://raw.githubusercontent.com/LidiaRipollSanchez/Neuropeptide-Connectome/6689619236ba1b4681a9a77b3d918d513416336c/'
FILES = {
    'nemanode_cells.json': 'https://nemanode.org/api/cells',
    'witvliet/nemanode/datasets.json': 'https://nemanode.org/api/datasets',
    'witvliet/41586_2021_3778_MOESM8_ESM.csv': NATURE.format('s41586-021-03778-8', '41586_2021_3778_MOESM8_ESM.csv'),
    **{f'witvliet/nemanode/witvliet_2020_{i}.json': f'https://nemanode.org/api/download-connectivity?datasetId=witvliet_2020_{i}'
       for i in range(1, 9)},
    'cook/moesm12.xlsx': NATURE.format('s41586-019-1352-7', '41586_2019_1352_MOESM12_ESM.xlsx'),
    'cook/moesm13.xlsx': NATURE.format('s41586-019-1352-7', '41586_2019_1352_MOESM13_ESM.xlsx'),
    'cook/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx':
        'https://wormwiring.org/si/SI%205%20Connectome%20adjacency%20matrices,%20corrected%20July%202020.xlsx',
    **{f'neuropeptide/01022024_neuropeptide_connectome_{k}_range_model.csv':
       NP + f'Adjacency%20matrices%20for%20networks/01022024_neuropeptide_connectome_{k}_range_model.csv'
       for k in ('short', 'mid', 'long')},
    'neuropeptide/072022_anatomical_class.csv': NP + 'Scripts%20%26%20data/072022_anatomical_class.csv',
}
CLASSES = ['stable', 'variable', 'developmentally dynamic (strengthened)', 'developmentally dynamic (weakened)',
           'post-embryonic brain integration']


def fetch(rel):
    p = CACHE / rel
    if not p.exists():
        p.parent.mkdir(parents=True, exist_ok=True)
        print('fetching', FILES[rel], file=sys.stderr)
        req = urllib.request.Request(FILES[rel], headers={'User-Agent': 'moonshine-fetch'})
        p.write_bytes(urllib.request.urlopen(req, timeout=120).read())
    return p


def write(name, obj):
    p = OUT / name
    p.write_text(json.dumps(obj, ensure_ascii=False, separators=(',', ':')))
    print(f'-> {p.relative_to(ROOT)} ({p.stat().st_size:,} B)', file=sys.stderr)


CELLS = None


def cell(n):  # NemaNode writes AS1 where WormWiring and the peptide matrices write AS01
    return CELLS.get(n) or CELLS[re.sub(r'0(\d)$', r'\1', n)]


def role(n):
    t = cell(n)['type']
    if t == 'b':
        return 'b'
    return next((ch for ch in t if ch in 'sim'), 'i' if 'n' in t else 'o')  # 'n' only: modulatory, as article 3


def witvliet():
    ages = {d['id']: d['time'] for d in json.loads(fetch('witvliet/nemanode/datasets.json').read_text())}
    cls = {(r['pre'], r['post']): r['classification'] for r in csv.DictReader(open(fetch('witvliet/41586_2021_3778_MOESM8_ESM.csv')))}
    ds = []
    for i in range(1, 9):
        rows = json.loads(fetch(f'witvliet/nemanode/witvliet_2020_{i}.json').read_text())
        ds.append({(r['pre'], r['post']): r['synapses'] for r in rows if r['type'] == 'chemical'})
    edges = sorted(set().union(*ds))
    assert set(edges) == set(cls), 'MOESM8 is the union of the eight chemical edge lists'
    names = sorted({n for e in edges for n in e})
    idx = {n: i for i, n in enumerate(names)}
    out = {
        'source': 'Witvliet et al. 2021, Nature 596: 257-261; edge lists via NemaNode (witvliet_2020_1..8), '
                  'connection classes from Supplementary Table 6; cell types from NemaNode /api/cells',
        'ages': [ages[f'witvliet_2020_{i}'] for i in range(1, 9)],
        'stages': ['L1', 'L1', 'L1', 'L1', 'L2', 'L3', 'adult', 'adult'],
        'classes': CLASSES,
        'nodes': [{'name': n, 'cls': cell(n)['class'], 'role': role(n)} for n in names],
        'chem': [[idx[a], idx[b], CLASSES.index(cls[(a, b)]), [d.get((a, b), 0) for d in ds]] for a, b in edges],
    }
    assert out['ages'] == [0, 5, 8, 16, 23, 27, 50, 50]
    print(f'Witvliet: {len(names)} cells, {len(edges)} chemical connections, synapses per dataset '
          f'{[sum(d.values()) for d in ds]}', file=sys.stderr)
    write('worm-witvliet.json', out)


def sheet(path, name):
    return list(openpyxl.load_workbook(path, read_only=True)[name].iter_rows(values_only=True))


def sexes():
    lrp, sxp = fetch('cook/moesm12.xlsx'), fetch('cook/moesm13.xlsx')
    lr = {k: [[r[0], r[1], r[2], r[3], round(r[4], 6)] for r in sheet(lrp, s)[1:] if r[0]]
          for k, s in (('chem', 'chem_lr_herm_stats'), ('gap', 'gap_jn_lr_herm_stats'))}
    sex = {}
    book = openpyxl.load_workbook(sxp)
    for k, alls, dirs in (('chem', 'chem_zscores_all_pairs', ('chem M&gt;H', 'chem H&gt;M')),
                          ('gap', 'gap_jn_zscores_all_pairs', ('gap jn M&gt;H', 'gap jn H&gt;M'))):
        flags = {}
        for s in dirs:  # the marks are cell fills under the flag columns, not values
            ws = book[s.replace('&gt;', '>')]
            head = [c.value for c in ws[2]]
            for r in ws.iter_rows(min_row=3):
                if not r[0].value:
                    continue
                marks = {head[c.column - 1] for c in r[8:] if c.fill is not None and c.fill.fill_type}
                f = 'c' if 'this work' in marks else 'o' if 'Oren-Suisse' in marks else 'd' if marks else ''
                key = (r[0].value, r[1].value)
                flags[key] = flags.get(key) or f  # AVG->VD13 and AVG->DA09 are listed twice
        # The gap-junction 'all pairs' sheet only holds the hermaphrodite-stronger rows (it equals
        # 'gap jn H>M'), so take the union of the all-pairs sheet and both direction sheets.
        # Three gap-junction pairs (ALM-PVR, ALA-RID, ALM-AVM) appear in both direction sheets
        # with contradictory weights; they are dropped.
        pairs, bad = {}, set()
        for s in (alls,) + tuple(d.replace('&gt;', '>') for d in dirs):
            for r in sheet(sxp, s):
                if not r[0] or r[0] == 'pre' or not isinstance(r[2], (int, float)):
                    continue
                v = (r[2], r[3], r[4], round(r[7], 6))
                if pairs.setdefault((r[0], r[1]), v) != v:
                    bad.add((r[0], r[1]))
        assert len(bad) == (3 if k == 'gap' else 0), bad
        sex[k] = [[a, b, *v, flags.get((a, b), '')] for (a, b), v in pairs.items() if (a, b) not in bad]
    print(f'Cook: left/right {len(lr["chem"])} chemical, {len(lr["gap"])} gap; sex {len(sex["chem"])} chemical, '
          f'{len(sex["gap"])} gap; flagged {sum(1 for r in sex["chem"] if r[6])} chemical', file=sys.stderr)
    assert len(lr['chem']) == 1285 and len(sex['chem']) == 1821, len(sex['chem'])
    write('worm-sexes.json', {
        'source': 'Cook et al. 2019, Nature 571: 63-71, Supplementary Information 8 (left/right) and 9 (sex differences); '
                  'weights are EM serial sections',
        'lr': lr, 'sex': sex})


def peptide():
    groups_rows = list(csv.DictReader(open(fetch('neuropeptide/072022_anatomical_class.csv'))))
    group = {r['Neuron'].strip("'"): r['Anatomical cell class (WW Barry)'] for r in groups_rows}
    mats = {}
    for k in ('short', 'mid', 'long'):
        rows = list(csv.reader(open(fetch(f'neuropeptide/01022024_neuropeptide_connectome_{k}_range_model.csv'))))
        names = rows[0][1:]
        assert [r[0] for r in rows[1:]] == names and len(names) == 302
        mats[k] = (names, [[int(float(x or 0)) for x in r[1:]] for r in rows[1:]])
    names = mats['short'][0]
    assert all(mats[k][0] == names for k in mats)
    ws = sheet(fetch('cook/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx'), 'hermaphrodite chemical')
    col = {c: j for j, c in enumerate(ws[2]) if c}
    row = {r[2]: r for r in ws[3:] if r[2]}
    idx = {n: i for i, n in enumerate(names)}
    syn = []
    for a in names:
        if a not in row:
            continue  # CANL, CANR make no chemical synapses
        for b in names:
            v = row[a][col[b]] if b in col else None
            if v and a != b:
                syn.append([idx[a], idx[b], int(v)])
    groups = sorted(set(group.values()), key=lambda g: names.index(next(n for n in names if group[n] == g)))
    enc = lambda M: ''.join('0123456789abcdefghijklmnopqrstuvwxyz'[v] for r in M for v in r)
    out = {
        'source': 'Ripoll-Sanchez et al. 2023, Neuron 111: 3570-3589, aggregate neuropeptide matrices (2024-02-02 update, '
                  'commit 6689619); synapses from Cook et al. 2019 SI 5 (hermaphrodite chemical, corrected July 2020)',
        'groups': groups,
        'nodes': [{'name': n, 'cls': cell(n)['class'], 'role': role(n), 'group': groups.index(group[n])} for n in names],
        'syn': syn,
        'pep': {k: enc(M) for k, (_, M) in mats.items()},
    }
    print(f'Peptides: {len(names)} neurons, {len(syn)} chemical pairs, nonzero '
          f'{ {k: sum(v > 0 for r in M for v in r) for k, (_, M) in mats.items()} }', file=sys.stderr)
    write('worm-peptide.json', out)


def main():
    global CELLS
    CELLS = {c['name']: c for c in json.loads(fetch('nemanode_cells.json').read_text())}
    witvliet()
    sexes()
    peptide()


if __name__ == '__main__':
    main()
